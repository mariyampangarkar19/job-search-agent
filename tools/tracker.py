"""Job application tracker (Excel mirror of the Google Sheet the agent keeps).

  python3 tracker.py create            # build a fresh Job_Tracker.xlsx (refuses to overwrite)
  python3 tracker.py add roles.json    # append new roles, skipping links already in the tracker
  python3 tracker.py expire 12 40      # set Status to "Link expired" for those IDs
  python3 tracker.py remove 12 40      # delete those rows entirely
  python3 tracker.py relink new.json   # replace links, from {"id": "new url"}

roles.json is a list of objects with keys: country, city, company, title, url, source,
posted, key_skills, sponsorship, fit_notes, and optionally fit (1-5) and priority (1-3).
"""
import json
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TRACKER = Path(__file__).resolve().parent.parent / "Job_Tracker.xlsx"
MAX_ROW = 2000

COLUMNS = [  # (header, width)
    ("ID", 6), ("Added", 11), ("Priority", 8), ("Fit (1-5)", 8), ("Country", 12), ("City", 12),
    ("Company", 20), ("Role", 38), ("Link", 12), ("Source", 13), ("Posted", 11),
    ("Resume (auto)", 16), ("Resume override", 14), ("Resume to send", 16),
    ("Tweak: add these JD keywords (only if true)", 34), ("Sponsorship", 28), ("Why it fits", 44),
    ("Status", 17), ("Date applied", 12), ("Follow-up due", 12), ("Contact name", 18),
    ("Contact LinkedIn", 18), ("Message sent?", 10), ("Notes", 30),
]
COL = {h: get_column_letter(i + 1) for i, (h, _) in enumerate(COLUMNS)}

STATUSES = ["To apply", "Applied", "Recruiter messaged", "Screening", "Technical interview",
            "Final round", "Offer", "Rejected", "No response", "Withdrawn", "Skipped", "Link expired"]
RESPONSES = ["Screening", "Technical interview", "Final round", "Offer"]

# First match wins, top to bottom. Keywords are matched case-insensitively against " <title> ".
RULES = [  # 2026-10-27: collapsed to AI Engineer / Software Engineer, with Technical Analyst for analyst and product roles
    ("AI Engineer", "resumes/<UAE, EU or India>/Mariyam_Pangarkar_AI_Engineer.pdf",
     ["forward deployed", "solutions engineer", "solution engineer", "customer engineer",
      "implementation engineer", "deployment strategist",
      " ai ", " ai/", "/ai ", " ai-", "(ai", "ai engineer", "genai", "gen ai", "generative",
      "applied ai", "machine learning", " ml ", " ml/", "/ml ", "llm", "agent", "prompt", "nlp"]),
    ("Technical Analyst", "resumes/<UAE, EU or India>/Mariyam_Pangarkar_Technical_Analyst.pdf",
     ["technical analyst", "business analyst", "systems analyst", "system analyst", "product owner",
      "technical product", "product manager", "solutions analyst", "integration analyst",
      "data analyst", "data scientist", "data science", " bi ", "business intelligence", "product analyst"]),
]
DEFAULT = ("Software Engineer", "resumes/<UAE, EU or India>/Mariyam_Pangarkar_Software_Engineer.pdf")

# Skills already on her resumes; JD skills outside this set become the tweak suggestion.
HER_SKILLS = {s.lower() for s in """python java typescript javascript node.js node c++ c# .net sql postgresql
postgres rest rest apis apis aws gcp docker jenkins github actions ci/cd cloudflare supabase linux bash dynamodb
opensearch pandas numpy etl data pipelines data migration data analysis unit testing integration testing pytest
postman testrail agile technical documentation generative ai genai langchain langgraph rag chromadb n8n react
next.js nextjs angular vite tailwind css webhooks websockets payment integrations payments aws lambda lambda
step functions api gateway iam flutter zustand tanstack query material ui tableau power bi eda statistics
hypothesis testing tensorflow django graphql llm llms git microservices blockchain solana fintech kyc
stakeholder management requirements sql databases system design""".replace("\n", " ").split(" ")} | {
    "rest apis", "github actions", "data pipelines", "next.js", "node.js", "aws lambda", "power bi",
    "generative ai", "system design", "stakeholder management", "unit testing", "integration testing"}

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
RED = PatternFill("solid", fgColor="F8D7DA")
GREEN = PatternFill("solid", fgColor="D4EDDA")
BLUE = PatternFill("solid", fgColor="DCEBFA")
GREY = PatternFill("solid", fgColor="E9ECEF")


def resume_formula(row):
    cell = f"{COL['Role']}{row}"
    expr = f'"{DEFAULT[0]}"'
    for label, _, words in reversed(RULES):
        tests = ",".join(f'ISNUMBER(SEARCH("{w}"," "&{cell}&" "))' for w in words)
        expr = f'IF(OR({tests}),"{label}",{expr})'
    return f'=IF({cell}="","",{expr})'


def tweak(key_skills):
    out = []
    for raw in (key_skills or "").split(","):
        s = raw.strip()
        if "/" in s or len(s.split()) > 3:
            continue
        if s and s.lower() not in HER_SKILLS and s.lower().rstrip("s") not in HER_SKILLS:
            out.append(s)
    return ", ".join(out[:3])


def stable_link(url):
    """Normalise Greenhouse embed links to the regular job page; embed links can redirect to the job board."""
    m = re.search(r"(job-boards(?:\.eu)?\.greenhouse\.io)/embed/job_app\?for=([\w-]+)&token=(\d+)", url)
    if m:
        return f"https://{m.group(1)}/{m.group(2)}/jobs/{m.group(3)}"
    return url


def write_row(ws, r, idx, job):
    c = lambda h: f"{COL[h]}{r}"
    ws[c("ID")] = idx
    ws[c("Added")] = date.today()
    ws[c("Added")].number_format = "yyyy-mm-dd"
    ws[c("Priority")] = job.get("priority", 2)
    ws[c("Fit (1-5)")] = job.get("fit", 3)
    for h, k in [("Country", "country"), ("City", "city"), ("Company", "company"), ("Role", "title"),
                 ("Source", "source"), ("Posted", "posted"), ("Sponsorship", "sponsorship"),
                 ("Why it fits", "fit_notes")]:
        ws[c(h)] = job.get(k, "")
    ws[c("Link")] = stable_link(job["url"])
    ws[c("Link")].hyperlink = stable_link(job["url"])
    ws[c("Link")].font = Font(color="0563C1", underline="single")
    ws[c("Resume (auto)")] = resume_formula(r)
    ws[c("Resume override")] = job.get("resume_override", "")
    ws[c("Resume to send")] = f'=IF({c("Resume override")}<>"",{c("Resume override")},{c("Resume (auto)")})'
    ws[c("Tweak: add these JD keywords (only if true)")] = job.get("tweak", tweak(job.get("key_skills")))
    ws[c("Status")] = "To apply"
    ws[c("Follow-up due")] = f'=IF({c("Date applied")}="","",WORKDAY({c("Date applied")},5))'
    ws[c("Follow-up due")].number_format = "yyyy-mm-dd"
    ws[c("Date applied")].number_format = "yyyy-mm-dd"
    for h in COL:
        ws[c(h)].alignment = Alignment(vertical="top", wrap_text=h in ("Role", "Why it fits", "Sponsorship",
                                                                          "Tweak: add these JD keywords (only if true)", "Notes"))


def create():
    if TRACKER.exists():
        sys.exit(f"{TRACKER} exists; not overwriting")
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    for i, (h, w) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "I2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{MAX_ROW}"

    rng = lambda h: f"{COL[h]}2:{COL[h]}{MAX_ROW}"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv.add(rng("Status"))
    yn = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    yn.add(rng("Message sent?"))
    labels = [r[0] for r in RULES] + [DEFAULT[0]]
    ov = DataValidation(type="list", formula1='"' + ",".join(labels) + '"', allow_blank=True)
    ov.add(rng("Resume override"))
    pr = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    pr.add(rng("Priority"))
    for v in (dv, yn, ov, pr):
        ws.add_data_validation(v)

    s, f = COL["Status"], COL["Follow-up due"]
    ws.conditional_formatting.add(
        rng("Follow-up due"),
        FormulaRule(formula=[f'AND(${f}2<>"",${f}2<=TODAY(),OR(${s}2="Applied",${s}2="Recruiter messaged"))'], fill=RED))
    full = f"A2:{get_column_letter(len(COLUMNS))}{MAX_ROW}"
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'${s}2="Offer"'], fill=GREEN))
    ws.conditional_formatting.add(full, FormulaRule(
        formula=[f'OR(${s}2="Screening",${s}2="Technical interview",${s}2="Final round")'], fill=BLUE))
    ws.conditional_formatting.add(full, FormulaRule(
        formula=[f'OR(${s}2="Rejected",${s}2="Withdrawn",${s}2="Skipped",${s}2="No response",${s}2="Link expired")'], fill=GREY))

    rules = wb.create_sheet("Resume Rules")
    rules.append(["Order", "Resume label", "File in the project folder", "Job title contains any of"])
    for i, (label, fname, words) in enumerate(RULES, 1):
        rules.append([i, label, fname, ", ".join(w.strip() for w in words)])
    rules.append([len(RULES) + 1, DEFAULT[0], DEFAULT[1], "anything else (Software Engineer, SDE, Backend, Platform)"])
    rules.append([])
    rules.append(["", "Override", "", "If a generic title hides a specialist JD (e.g. 'Software Engineer' at an AI startup), "
                  "pick a label in 'Resume override' on the Applications sheet."])
    for i, w in enumerate([7, 18, 40, 90], 1):
        rules.column_dimensions[get_column_letter(i)].width = w
    for cell in rules[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL

    dash = wb.create_sheet("Dashboard")
    A = "Applications!"
    R = lambda h: f"{A}${COL[h]}$2:${COL[h]}${MAX_ROW}"
    applied = lambda extra="": (f'SUMPRODUCT({extra}({R("Status")}<>"")*({R("Status")}<>"To apply")'
                                f'*({R("Status")}<>"Skipped")*({R("Status")}<>"Link expired"))')
    resp = lambda extra="": ("SUMPRODUCT(" + extra + "((" + "+".join(
        f'({R("Status")}="{x}")' for x in RESPONSES) + ")>0))")
    dash.append(["Overall", "Count"])
    dash.append(["Roles in tracker", f'=COUNTA({R("Role")})'])
    dash.append(["Waiting to apply", f'=COUNTIF({R("Status")},"To apply")'])
    dash.append(["Applied (any stage after To apply)", "=" + applied()])
    dash.append(["Got a response (screening or later)", "=" + resp()])
    dash.append(["Response rate", "=IFERROR(B5/B4,0)"])
    dash.append(["Follow-ups due today or overdue",
                 f'=SUMPRODUCT(({R("Follow-up due")}<>"")*({R("Follow-up due")}<=TODAY())'
                 f'*((({R("Status")}="Applied")+({R("Status")}="Recruiter messaged"))>0))'])
    dash.append(["Offers", f'=COUNTIF({R("Status")},"Offer")'])
    dash["B6"].number_format = "0%"
    dash.append([None])
    dash.append(["By market", "Applied", "Responses", "Response rate"])
    start = dash.max_row
    for i, country in enumerate(["UAE", "Germany", "Netherlands", "Ireland", "Luxembourg", "India"]):
        r = start + 1 + i
        m = f'({R("Country")}=A{r})*'
        dash.append([country, "=" + applied(m), "=" + resp(m), f"=IFERROR(C{r}/B{r},0)"])
        dash[f"D{r}"].number_format = "0%"
    dash.append([None])
    dash.append(["By resume", "Applied", "Responses", "Response rate"])
    start = dash.max_row
    for i, label in enumerate([r[0] for r in RULES] + [DEFAULT[0]]):
        r = start + 1 + i
        m = f'({R("Resume to send")}=A{r})*'
        dash.append([label, "=" + applied(m), "=" + resp(m), f"=IFERROR(C{r}/B{r},0)"])
        dash[f"D{r}"].number_format = "0%"
    for row in dash.iter_rows():
        if row[0].value in ("Overall", "By market", "By resume"):
            for cell in row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = HEADER_FILL
    dash.column_dimensions["A"].width = 38
    for col in "BCD":
        dash.column_dimensions[col].width = 14

    how = wb.create_sheet("How to use", 0)
    for line in [
        "How to use this tracker",
        "",
        "1. Applications tab: sort or filter by Priority (1 first), then Fit.",
        "2. 'Resume to send' tells you which PDF to attach (file names are on the Resume Rules tab). Never paste a JD to ask.",
        "3. If the Tweak column lists keywords, add them to your Skills line only if you have really used them.",
        "4. After applying: set Status to Applied and type the Date applied. Follow-up due fills in (5 working days) and turns red when due.",
        "5. Message the recruiter or hiring manager the same day; put their name in Contact name and set Message sent? to Yes.",
        "6. Move Status as things progress. Rows turn blue for interviews, green for an offer, grey when closed.",
        "7. Dashboard tab: response rate by market and by resume. Claude reads this every Friday to adjust the next week.",
        "",
        "Claude adds new roles to the bottom of the Applications tab each weekday. Don't rename or reorder the columns.",
    ]:
        how.append([line])
    how["A1"].font = Font(bold=True, size=14)
    how.column_dimensions["A"].width = 130
    wb.active = wb.index(ws)
    wb.save(TRACKER)
    print(f"created {TRACKER}")


def add(path):
    jobs = json.loads(Path(path).read_text())
    wb = load_workbook(TRACKER)
    ws = wb["Applications"]
    link_col = COLUMNS.index(("Link", 12)) + 1
    seen, last, max_id = set(), 1, 0
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=link_col)
        if cell.hyperlink is not None:
            seen.add(cell.hyperlink.target)
        if ws.cell(row=r, column=8).value:
            last = r
            max_id = max(max_id, int(ws.cell(row=r, column=1).value or 0))
    keyset = {(str(ws.cell(row=r, column=7).value).lower(), str(ws.cell(row=r, column=8).value).lower())
              for r in range(2, last + 1)}
    added = 0
    for job in jobs:
        if stable_link(job["url"]) in seen or job["url"] in seen or (job["company"].lower(), job["title"].lower()) in keyset:
            continue
        last += 1
        max_id += 1
        write_row(ws, last, max_id, job)
        seen.add(job["url"])
        keyset.add((job["company"].lower(), job["title"].lower()))
        added += 1
    wb.save(TRACKER)
    print(f"added {added} new roles ({len(jobs) - added} duplicates skipped); tracker has {last - 1} rows")


def expire(ids):
    """Set Status to 'Link expired' for the given tracker IDs."""
    wb = load_workbook(TRACKER)
    ws = wb["Applications"]
    status_col = list(COL).index("Status") + 1
    hit = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value in ids:
            ws.cell(row=r, column=status_col).value = "Link expired"
            hit += 1
    wb.save(TRACKER)
    print(f"marked {hit} rows Link expired")


def _rewrite(ws, rows):
    """Write rows (lists of (value, hyperlink) per column) from row 2 down, regenerating per-row formulas."""
    formula_cols = {"Resume (auto)", "Resume to send", "Follow-up due"}
    for r in range(2, ws.max_row + 1):
        for col in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=col)
            cell.value, cell.hyperlink = None, None
    for r, row in enumerate(rows, start=2):
        c = lambda h: f"{COL[h]}{r}"
        for (h, _), (value, link) in zip(COLUMNS, row):
            if h not in formula_cols:
                ws[c(h)] = value
        link = row[list(COL).index("Link")][1]
        ws[c("Link")].hyperlink = link
        ws[c("Link")].font = Font(color="0563C1", underline="single")
        ws[c("Resume (auto)")] = resume_formula(r)
        ws[c("Resume to send")] = f'=IF({c("Resume override")}<>"",{c("Resume override")},{c("Resume (auto)")})'
        ws[c("Follow-up due")] = f'=IF({c("Date applied")}="","",WORKDAY({c("Date applied")},5))'
        for h in ("Added", "Date applied", "Follow-up due"):
            ws[c(h)].number_format = "yyyy-mm-dd"


def _read_rows(ws):
    rows = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is None:
            continue
        rows.append([(ws.cell(row=r, column=col).value,
                      ws.cell(row=r, column=col).hyperlink.target if ws.cell(row=r, column=col).hyperlink else None)
                     for col in range(1, len(COLUMNS) + 1)])
    return rows


def remove(ids):
    """Delete the rows with these tracker IDs; IDs of the remaining rows are kept."""
    wb = load_workbook(TRACKER)
    ws = wb["Applications"]
    rows = _read_rows(ws)
    kept = [row for row in rows if row[0][0] not in ids]
    _rewrite(ws, kept)
    wb.save(TRACKER)
    print(f"removed {len(rows) - len(kept)} rows; tracker has {len(kept)} rows")


def relink(path):
    """Replace links from a JSON object {id: new_url}."""
    new = {int(k): stable_link(v) for k, v in json.loads(Path(path).read_text()).items()}
    wb = load_workbook(TRACKER)
    ws = wb["Applications"]
    link_col = list(COL).index("Link") + 1
    for r in range(2, ws.max_row + 1):
        i = ws.cell(row=r, column=1).value
        if i in new:
            ws.cell(row=r, column=link_col).value = new[i]
            ws.cell(row=r, column=link_col).hyperlink = new[i]
    wb.save(TRACKER)
    print(f"relinked {len(new)} rows")


if __name__ == "__main__":
    if sys.argv[1:2] == ["create"]:
        create()
    elif sys.argv[1:2] == ["add"] and len(sys.argv) == 3:
        add(sys.argv[2])
    elif sys.argv[1:2] == ["expire"] and len(sys.argv) > 2:
        expire({int(x) for x in sys.argv[2:]})
    elif sys.argv[1:2] == ["remove"] and len(sys.argv) > 2:
        remove({int(x) for x in sys.argv[2:]})
    elif sys.argv[1:2] == ["relink"] and len(sys.argv) == 3:
        relink(sys.argv[2])
    else:
        sys.exit(__doc__)
