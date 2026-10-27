"""Print tracker rows as Google Sheet rows (Applications tab layout since 2026-10-28).

Columns: ID | Country | City | Company | Role | Link | Resume to send | Status | Date applied | Follow-up due | Notes
Usage: python3 tools/sheet_rows.py <id> [<id> ...]   -> JSON list of rows (Follow-up due left blank; set it with
update_formulas  =IF(I{r}="","",WORKDAY(I{r},5))  after appending).
"""
import json, sys
import openpyxl
from tracker import RULES, DEFAULT, TRACKER


def pick_resume(title):
    t = f' {title.lower()} '
    for label, _, words in RULES:
        if any(w in t for w in words):
            return label
    return DEFAULT[0]

ws = openpyxl.load_workbook(str(TRACKER))['Applications']
ids = {int(x) for x in sys.argv[1:]}
rows = []
for r in range(2, ws.max_row + 1):
    i = ws.cell(r, 1).value
    if i not in ids:
        continue
    resume = ws.cell(r, 13).value or ws.cell(r, 12).value  # override, else auto
    if not resume or (isinstance(resume, str) and resume.startswith('=')):
        resume = pick_resume(ws.cell(r, 8).value or '')
    rows.append([i, ws.cell(r, 5).value, ws.cell(r, 6).value, ws.cell(r, 7).value, ws.cell(r, 8).value,
                 ws.cell(r, 9).value, resume or '', ws.cell(r, 18).value or 'To apply', '', '', ws.cell(r, 24).value or ''])
print(json.dumps(rows, default=str))
